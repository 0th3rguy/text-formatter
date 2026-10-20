# 检验「重新排版」产出是否真的把图片带进了新文档。
import sys, zipfile, os, re
from docx import Document

path = sys.argv[1] if len(sys.argv) > 1 else r"D:\WorkSpace\text-formatter\test\out\reflowed-with-images.docx"
print("文件:", path, os.path.getsize(path), "字节")

z = zipfile.ZipFile(path)
print("ZIP 完整性:", "OK" if z.testzip() is None else "损坏")
media = [n for n in z.namelist() if n.startswith("word/media/")]
print("内嵌媒体:", media)
for m in media:
    print("   ", m, len(z.read(m)), "字节")

doc_xml = z.read("word/document.xml").decode("utf-8", "ignore")
print("document.xml 中 <w:drawing> 数:", doc_xml.count("<w:drawing>"))
print("document.xml 中 a:blip 数:", len(re.findall(r"<a:blip[^>]*r:embed=", doc_xml)))
print("document.xml 中 r:embed 引用:", sorted(set(re.findall(r'r:embed="([^"]+)"', doc_xml))))

rels = z.read("word/_rels/document.xml.rels").decode("utf-8", "ignore")
img_rels = re.findall(r'Id="([^"]+)"[^>]*Type="[^"]*image"[^>]*Target="([^"]+)"', rels)
print("document.xml.rels 图片关系:", img_rels)

ct = z.read("[Content_Types].xml").decode("utf-8", "ignore")
print("Content_Types 中的图片扩展名声明:", re.findall(r'Extension="(png|jpe?g|gif|bmp|webp)"', ct))

d = Document(path)
print("段落数:", len(d.paragraphs), " 表格数:", len(d.tables))
# 每张图的实际尺寸
import docx
for i, sh in enumerate(d.inline_shapes, 1):
    print("  图%d: %.2f x %.2f cm" % (i, sh.width.cm, sh.height.cm))

ok = len(media) == 2 and doc_xml.count("<w:drawing>") == 2 and len(img_rels) == 2
print("\n结论:", "图片完整保留 ✔" if ok else "图片有缺失 ✘")
sys.exit(0 if ok else 1)
