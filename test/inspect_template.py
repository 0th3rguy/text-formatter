# 逐段 dump 一个 .docx 的真实排版信息（字体/字号/缩进/行距/对齐），作为"正规模板"的基准。
import sys, os, zipfile, re
from docx import Document
from docx.shared import Pt

DEFAULT = r"D:\Work\政务信息\2026\2026.6.23卷烟物流配送中心“四个强化”夯实运营管理根基.docx"
path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT
print("文件:", path)
print("存在:", os.path.exists(path), " 大小:", os.path.getsize(path) if os.path.exists(path) else "-")
print("=" * 78)

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

doc = Document(path)

# ---- 节 / 页面 ----
for i, s in enumerate(doc.sections):
    print("节%d: 纸 %.2fx%.2fcm 边距 上%.2f 下%.2f 左%.2f 右%.2f 页眉距%.2f 页脚距%.2f" % (
        i, s.page_width.cm, s.page_height.cm, s.top_margin.cm, s.bottom_margin.cm,
        s.left_margin.cm, s.right_margin.cm, s.header_distance.cm, s.footer_distance.cm))
print("=" * 78)

# ---- styles.xml 里的默认字体 ----
z = zipfile.ZipFile(path)
st = z.read("word/styles.xml").decode("utf-8", "ignore")
m = re.search(r"<w:docDefaults>.*?</w:docDefaults>", st, re.S)
if m:
    dd = m.group(0)
    rf = re.findall(r'w:eastAsia="([^"]*)"', dd)
    asc = re.findall(r'w:ascii="([^"]*)"', dd)
    sz = re.findall(r'<w:sz w:val="(\d+)"', dd)
    print("样式默认: eastAsia=%s ascii=%s sz(half-pt)=%s" % (rf[:3], asc[:3], sz[:3]))
print("=" * 78)

print("%-4s %-14s %-14s %-6s %-5s %-10s %-22s %s" % (
    "#", "样式", "中文字体", "字号", "粗", "对齐", "首行缩进", "行距/段距"))
print("-" * 110)

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if not t and i > 12:
        continue
    pf = p.paragraph_format
    # 直接格式里的字体
    ea = asc = None
    for r in p.runs:
        rpr = r._element.find(W + "rPr")
        if rpr is not None:
            rf = rpr.find(W + "rFonts")
            if rf is not None:
                ea = rf.get(W + "eastAsia") or ea
                asc = rf.get(W + "ascii") or asc
        if ea:
            break
    # 段落级 rPr 字体
    ppr = p._element.find(W + "pPr")
    pe = pa = None
    if ppr is not None:
        prpr = ppr.find(W + "rPr")
        if prpr is not None:
            rf = prpr.find(W + "rFonts")
            if rf is not None:
                pe = rf.get(W + "eastAsia")
                pa = rf.get(W + "ascii")
    # 缩进原始 XML（首行缩进的字符表示在这里）
    ind_xml = ""
    if ppr is not None:
        ind = ppr.find(W + "ind")
        if ind is not None:
            ind_xml = " ".join("%s=%s" % (k.split("}")[-1], v) for k, v in ind.attrib.items())
    szv = None
    for r in p.runs:
        if r.font.size:
            szv = r.font.size.pt
            break
    script = (pe or "—") + "/" + (ea or "—") if (pe or ea) else "—"
    print("%-4d %-14s %-14s %-6s %-5s %-10s %-22s" % (
        i, (p.style.name or "?")[:13], script[:13],
        ("%.1f" % szv) if szv else "-",
        "Y" if (p.runs and p.runs[0].bold) else ".",
        str(pf.alignment),
        ind_xml or "—"), end="")
    print(" 行距=%s%s 段前=%s 段后=%s" % (
        pf.line_spacing, "/" + str(pf.line_spacing_rule) if pf.line_spacing_rule else "",
        pf.space_before, pf.space_after))

print("=" * 78)
print("表格数:", len(doc.tables))
print("\n--- document.xml 中所有 w:ind 片段（去重） ---")
dx = z.read("word/document.xml").decode("utf-8", "ignore")
for frag in sorted(set(re.findall(r"<w:ind [^/]*/>", dx))):
    print("   ", frag)
print("\n--- 出现的 eastAsia 字体（去重计数） ---")
from collections import Counter
c = Counter(re.findall(r'w:eastAsia="([^"]*)"', dx))
for k, v in c.most_common():
    print("    %-24s %d" % (k, v))
print("\n--- 出现的字号 w:sz（去重计数） ---")
c2 = Counter(re.findall(r'<w:sz w:val="(\d+)"', dx))
for k, v in c2.most_common():
    print("    %s half-pt = %.1f pt   x%d" % (k, int(k) / 2, v))
