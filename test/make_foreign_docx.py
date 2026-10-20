# 造两个「外来」.docx 作为回读测试输入：
#   1. word-made.docx  —— 用真实 Microsoft Word COM 生成并保存（最权威的“别人给的文档”）
#   2. pydocx-made.docx —— 用 python-docx 生成（模拟程序导出的文档）
# 二者都不经过本项目代码，用来检验「重新排版已有文档」这条链路。
import os
import sys

OUT = r"D:\WorkSpace\text-formatter\test\out"
os.makedirs(OUT, exist_ok=True)

# ---------- 1. python-docx 生成 ----------
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

d = Document()
sec = d.sections[0]
sec.top_margin = Cm(2.54)
sec.left_margin = Cm(3.17)

t = d.add_heading("全市政务信息化建设情况报告", level=0)   # Title 样式
d.add_paragraph("—— 2026 年度工作总结").alignment = WD_ALIGN_PARAGRAPH.CENTER

d.add_paragraph("市政府办公室：")
d.add_paragraph("根据工作安排，现将全市政务信息化建设情况报告如下。")

d.add_heading("一、基本情况", level=1)
d.add_paragraph("全年共建成各类信息系统 37 个，覆盖政务服务、城市管理、民生保障等领域。")
d.add_heading("（一）基础设施", level=2)
d.add_paragraph("完成市级数据中心扩容，机架规模达到 1200 个标准机架。")
d.add_paragraph("1. 网络层面。骨干网带宽提升至 100G。")
d.add_paragraph("2. 算力层面。新增 GPU 算力 500 PFLOPS。")
d.add_heading("二、存在问题", level=1)

table = d.add_table(rows=3, cols=3)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "序号"; hdr[1].text = "问题"; hdr[2].text = "责任部门"
row = table.rows[1].cells
row[0].text = "1"; row[1].text = "数据共享不畅"; row[2].text = "大数据局"
row = table.rows[2].cells
row[0].text = "2"; row[1].text = "重复建设"; row[2].text = "发改委"

d.add_paragraph("三、下一步打算")
d.add_paragraph("持续推进数据资源整合，深化业务协同。")
d.add_paragraph("附件：问题清单")

p = os.path.join(OUT, "pydocx-made.docx")
d.save(p)
print("已生成:", p, os.path.getsize(p), "字节")

# ---------- 2. 真实 Word 生成 ----------
try:
    import win32com.client as win32
except ImportError:
    print("SKIP: 无 pywin32（Word 生成跳过）")
    sys.exit(0)

word = win32.Dispatch("Word.Application")
word.Visible = False
word.DisplayAlerts = 0
try:
    doc = word.Documents.Add()
    rng = doc.Content

    def add(text, style=None, align=None, size=None, bold=None):
        rng.InsertAfter(text + "\r")
        rng2 = doc.Paragraphs(doc.Paragraphs.Count).Range
        if style:
            try:
                rng2.Style = doc.Styles(style)
            except Exception as e:
                print("  样式设置失败", style, e)
        if align is not None:
            rng2.ParagraphFormat.Alignment = align
        if size:
            rng2.Font.Size = size
        if bold is not None:
            rng2.Font.Bold = bold
        return rng2

    # Word 常量：0=左 1=居中；样式名用中文界面下的内置名
    add("关于加强安全生产工作的通知", "标题", 1, 22)
    add("（安委办〔2026〕15 号）", None, 1, 16)
    add("各区县人民政府，市直各部门：")
    add("为切实做好当前安全生产工作，现就有关事项通知如下。")
    add("一、提高思想认识", "标题 1", 0, 16)
    add("各单位要充分认识安全生产工作的极端重要性，压实主体责任。")
    add("二、突出工作重点", "标题 1", 0, 16)
    add("（一）危险化学品", "标题 2", 0, 16)
    add("开展全覆盖排查，建立隐患台账。")

    t = doc.Tables.Add(doc.Paragraphs(doc.Paragraphs.Count).Range, 3, 3)
    t.Borders.Enable = True
    t.Cell(1, 1).Range.Text = "序号"
    t.Cell(1, 2).Range.Text = "重点领域"
    t.Cell(1, 3).Range.Text = "牵头单位"
    t.Cell(2, 1).Range.Text = "1"
    t.Cell(2, 2).Range.Text = "危化品"
    t.Cell(2, 3).Range.Text = "应急局"
    t.Cell(3, 1).Range.Text = "2"
    t.Cell(3, 2).Range.Text = "建筑施工"
    t.Cell(3, 3).Range.Text = "住建局"

    doc.Content.InsertAfter("\r三、强化责任落实\r")
    doc.Content.InsertAfter("对工作不力的，严肃追责问责。\r")
    doc.Content.InsertAfter("附件：重点领域分工表\r")

    p2 = os.path.join(OUT, "word-made.docx")
    doc.SaveAs2(p2, FileFormat=16)   # wdFormatDocumentDefault = .docx
    doc.Close(0)
    print("已生成:", p2, os.path.getsize(p2), "字节")
finally:
    word.Quit()
